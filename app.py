# -*- coding: utf-8 -*-
"""
投稿排版助手 v1.0.14
"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import os
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
import time
import platform
from cryptography.fernet import Fernet
import threading

CONFIG_DIR = Path.home() / ".tougao_assistant"
CONFIG_FILE = CONFIG_DIR / "config.enc"
KEY_FILE = CONFIG_DIR / "key.bin"
CONFIG_DIR.mkdir(exist_ok=True)

def center_window(win, width=None, height=None):
    """把窗口放到当前屏幕正中。"""
    win.update_idletasks()

    if width is None:
        width = win.winfo_width()
    if height is None:
        height = win.winfo_height()

    screen_width = win.winfo_screenwidth()
    screen_height = win.winfo_screenheight()
    x = max((screen_width - width) // 2, 0)
    y = max((screen_height - height) // 2, 0)
    win.geometry(f"{width}x{height}+{x}+{y}")


def show_alert(parent, title, message, button_text="确定 [O]", width=520, height=220):
    """可靠的模态提示框：按钮、Alt+字母、Enter、Esc、右上角关闭均可退出。"""
    win = tk.Toplevel(parent)
    win.title(title)
    win.transient(parent)
    win.resizable(False, False)

    frame = ttk.Frame(win, padding=18)
    frame.pack(fill="both", expand=True)

    ttk.Label(
        frame,
        text=message,
        justify="left",
        wraplength=width - 50
    ).pack(anchor="w", fill="both", expand=True, pady=(0, 14))

    def close_alert(event=None):
        # 先解除 grab，再销毁窗口，避免嵌套模态窗口时父窗口一直被锁住。
        try:
            win.grab_release()
        except tk.TclError:
            pass
        if win.winfo_exists():
            win.destroy()
        return "break"

    btn = ttk.Button(frame, text=button_text, command=close_alert)
    btn.pack(side="right")

    # 快捷键绑定在窗口本身，而不是按钮本身：
    # 即使焦点在其它控件上，Alt+字母 / Enter / Esc 仍然有效。
    m = re.search(r"\[([A-Za-z])\]", button_text)
    if m:
        key = m.group(1).lower()
        win.bind(f"<Alt-KeyPress-{key}>", close_alert)
    win.bind("<Return>", close_alert)
    win.bind("<Escape>", close_alert)
    win.protocol("WM_DELETE_WINDOW", close_alert)

    center_window(win, width, height)
    win.grab_set()
    win.focus_force()
    btn.focus_set()

    parent.wait_window(win)


def ask_retry_cancel(parent, title, message):
    """可靠的“重试/取消”模态对话框，打开后立即接管键盘焦点。"""
    result = {"value": False}

    win = tk.Toplevel(parent)
    win.title(title)
    win.transient(parent)
    win.resizable(False, False)

    frame = ttk.Frame(win, padding=18)
    frame.pack(fill="both", expand=True)

    ttk.Label(frame, text=message, justify="left", wraplength=500).pack(
        anchor="w", fill="both", expand=True, pady=(0, 14)
    )

    def choose(value, event=None):
        result["value"] = value
        try:
            win.grab_release()
        except tk.TclError:
            pass
        if win.winfo_exists():
            win.destroy()
        return "break"

    btn_retry = ttk.Button(frame, text="重试 [R]", command=lambda: choose(True))
    btn_cancel = ttk.Button(frame, text="取消 [C]", command=lambda: choose(False))
    btn_cancel.pack(side="right")
    btn_retry.pack(side="right", padx=(0, 8))

    win.bind("<Alt-KeyPress-r>", choose_retry := lambda e: choose(True, e))
    win.bind("<Alt-KeyPress-c>", choose_cancel := lambda e: choose(False, e))
    win.bind("<Return>", choose_retry)
    win.bind("<Escape>", choose_cancel)
    win.protocol("WM_DELETE_WINDOW", choose_cancel)

    center_window(win, 560, 220)
    win.grab_set()
    win.focus_force()
    btn_retry.focus_set()
    parent.wait_window(win)
    return result["value"]


def ask_export_conflict(parent, out_folder):
    """目标文件夹已存在时，返回 overwrite / saveas / cancel。"""
    result = {"value": None}

    win = tk.Toplevel(parent)
    win.title("目标文件夹已存在")
    win.transient(parent)
    win.resizable(False, False)

    frame = ttk.Frame(win, padding=18)
    frame.pack(fill="both", expand=True)

    ttk.Label(
        frame,
        text=f"文件夹：\n{out_folder}\n已经存在！",
        justify="left"
    ).pack(anchor="w", pady=(0, 12))

    ttk.Label(
        frame,
        text="请选择操作：\n"
             "覆盖：覆盖本次导出的同名文件（保留文件夹中的其他文件）\n"
             "另存为：给即将保存的文件夹改一个名字，仍保存在当前选定的目录下\n"
             "取消：放弃本次导出",
        justify="left"
    ).pack(anchor="w", pady=(0, 12))

    button_frame = ttk.Frame(frame)
    button_frame.pack(fill="x")

    def choose(value, event=None):
        result["value"] = value
        try:
            win.grab_release()
        except tk.TclError:
            pass
        if win.winfo_exists():
            win.destroy()
        return "break"

    btn_overwrite = ttk.Button(button_frame, text="覆盖 [O]", command=lambda: choose("overwrite"))
    btn_saveas = ttk.Button(button_frame, text="另存为 [S]", command=lambda: choose("saveas"))
    btn_cancel = ttk.Button(button_frame, text="取消 [C]", command=lambda: choose("cancel"))

    button_frame.columnconfigure(0, weight=1)
    btn_overwrite.grid(row=0, column=0, sticky="e", padx=(0, 8))
    btn_saveas.grid(row=0, column=1, sticky="e", padx=8)
    btn_cancel.grid(row=0, column=2, sticky="e", padx=(8, 0))

    win.bind("<Alt-KeyPress-o>", lambda e: choose("overwrite", e))
    win.bind("<Alt-KeyPress-s>", lambda e: choose("saveas", e))
    win.bind("<Alt-KeyPress-c>", lambda e: choose("cancel", e))
    win.bind("<Return>", lambda e: choose("overwrite", e))
    win.bind("<Escape>", lambda e: choose("cancel", e))
    win.protocol("WM_DELETE_WINDOW", lambda: choose("cancel"))

    center_window(win, 760, 230)
    win.grab_set()
    win.focus_force()
    btn_overwrite.focus_set()
    parent.wait_window(win)
    return result["value"]


def ask_saveas_folder_name(parent, initial_name):
    """让用户输入另存为的文件夹名称，并返回名称；取消时返回 None。"""
    result = {"value": None}

    win = tk.Toplevel(parent)
    win.title("另存为")
    win.transient(parent)
    win.resizable(False, False)

    frame = ttk.Frame(win, padding=18)
    frame.pack(fill="both", expand=True)

    ttk.Label(frame, text="请输入新的文件夹名称：", justify="left").pack(
        anchor="w", pady=(0, 8)
    )

    name_var = tk.StringVar(value=initial_name)
    entry = ttk.Entry(frame, textvariable=name_var, width=80)
    entry.pack(fill="x", pady=(0, 14))

    button_frame = ttk.Frame(frame)
    button_frame.pack(fill="x")

    def choose(value, event=None):
        if value == "ok":
            result["value"] = name_var.get()
        try:
            win.grab_release()
        except tk.TclError:
            pass
        if win.winfo_exists():
            win.destroy()
        return "break"

    btn_ok = ttk.Button(button_frame, text="确定 [O]", command=lambda: choose("ok"))
    btn_cancel = ttk.Button(button_frame, text="取消 [C]", command=lambda: choose("cancel"))
    btn_cancel.pack(side="right")
    btn_ok.pack(side="right", padx=(8, 0))

    win.bind("<Alt-KeyPress-o>", lambda e: choose("ok", e))
    win.bind("<Alt-KeyPress-c>", lambda e: choose("cancel", e))
    win.bind("<Return>", lambda e: choose("ok", e))
    win.bind("<Escape>", lambda e: choose("cancel", e))
    win.protocol("WM_DELETE_WINDOW", lambda: choose("cancel"))

    center_window(win, 720, 150)
    win.grab_set()
    win.focus_force()
    entry.focus_set()
    entry.selection_range(0, tk.END)
    parent.wait_window(win)
    return result["value"]


def get_cipher():
    if not KEY_FILE.exists():
        key = Fernet.generate_key()
        KEY_FILE.write_bytes(key)
    else:
        key = KEY_FILE.read_bytes()
    return Fernet(key)

def load_config():
    cipher = get_cipher()
    default_cfg = {
        "contact": {"姓名":"","微信":"","电话":"","地址":"","邮编":"","身份证号":"","邮箱":""},
        "banks": [{"银行名称":"","账号":"","开户支行":"","联行号":""}],
        "last_file_dir": "",
        "last_save_dir": ""
    }
    if not CONFIG_FILE.exists():
        return default_cfg
    try:
        raw = cipher.decrypt(CONFIG_FILE.read_bytes())
        cfg = json.loads(raw)
        for b in cfg.get("banks", []):
            if "卡号" in b and "账号" not in b:
                b["账号"] = b.pop("卡号")
        return cfg
    except Exception:
        return default_cfg

def save_config(cfg):
    cipher = get_cipher()
    raw = json.dumps(cfg, ensure_ascii=False, indent=2).encode("utf-8")
    CONFIG_FILE.write_bytes(cipher.encrypt(raw))

def _set_para(p, align=None, first_indent=None):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = first_indent

def _set_run_font(run, size_pt):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size_pt)

def convert_halfwidth_punct(s: str) -> str:
    """
    正文排版：将常见半角标点转换为全角标点，并自动配对中英文引号；书名号《》保持不变。
    映射：
    , . ! ? : ; () [] {} → ，。！？：；（）［］｛｝
    半角 " ' 做左右引号配对：奇数次左引号“ ‘，偶数次右引号” ’
    """
    trans_map = {
        ',': '，',
        '.': '。',
        '!': '！',
        '?': '？',
        ':': '：',
        ';': '；',
        '(': '（',
        ')': '）',
        '[': '［',
        ']': '］',
        '{': '｛',
        '}': '｝'
    }
    for half, full in trans_map.items():
        s = s.replace(half, full)

    # 处理双引号 "
    dq_count = 0
    out_chars = []
    for ch in s:
        if ch == '"':
            dq_count +=1
            if dq_count %2 ==1:
                out_chars.append("“")
            else:
                out_chars.append("”")
        else:
            out_chars.append(ch)
    s = "".join(out_chars)

    # 处理单引号 '
    sq_count =0
    out_chars2 = []
    for ch in s:
        if ch == "'":
            sq_count +=1
            if sq_count %2 ==1:
                out_chars2.append("‘")
            else:
                out_chars2.append("’")
        else:
            out_chars2.append(ch)
    return "".join(out_chars2)


def normalize_title_for_export(title):
    """
    根据标题中的《》/<>，生成：
    1. display_title：Word/TXT 文章中实际显示的标题
    2. file_title：文件夹名/文件名内部使用的标题

    Windows 文件名不能使用半角 < >，因此 file_title 中使用全角 ＜ ＞。
    """
    title = title.strip()

    # 如果《》或<>完整包住整个标题：排版时去掉最外层。
    if len(title) >= 2 and title.startswith("《") and title.endswith("》"):
        inner = title[1:-1].strip()
        if inner:
            return inner, make_windows_safe_filename(inner)

    if len(title) >= 2 and title.startswith("<") and title.endswith(">"):
        inner = title[1:-1].strip()
        if inner:
            return inner, make_windows_safe_filename(inner)

    def contains_chinese(s):
        return any("\u4e00" <= ch <= "\u9fff" for ch in s)

    import re
    pattern = re.compile(r"(《|<)([^《》<>]+)(》|>)")

    display_parts = []
    file_parts = []
    last = 0

    for match in pattern.finditer(title):
        display_parts.append(title[last:match.start()])
        file_parts.append(title[last:match.start()])

        left = match.group(1)
        content = match.group(2)

        # 文件夹/文件名内部统一使用 <内容> 的逻辑形式。
        file_marked = f"<{content}>"

        # 《》/<> 包裹中文或中英混合：文章标题使用《》。
        # 包裹全英文：文章标题使用<>。
        if contains_chinese(content):
            display_marked = f"《{content}》"
        else:
            display_marked = file_marked

        display_parts.append(display_marked)
        file_parts.append(file_marked)
        last = match.end()

    display_parts.append(title[last:])
    file_parts.append(title[last:])

    display_title = "".join(display_parts)
    file_title = make_windows_safe_filename("".join(file_parts))

    return display_title, file_title


def make_windows_safe_filename(name):
    """
    把 Windows 文件名禁止字符替换成视觉接近、但允许使用的全角字符。
    文章标题本身不修改，只用于文件夹名/文件名。
    """
    replacements = {
        "<": "＜",
        ">": "＞",
        ":": "：",
        '"': "＂",
        "/": "／",
        "\\": "＼",
        "|": "｜",
        "?": "？",
        "*": "＊",
    }

    safe = "".join(replacements.get(ch, ch) for ch in name)

    # Windows 不允许文件名/文件夹名以空格或句号结尾。
    safe = safe.rstrip(" .")
    return safe


def count_cn_and_word(body_lines):
    """
    按 WPS 常见“字数”口径统计正文：
    - 连续的 ASCII 字母/数字串作为 1 个单词，例如 V2、100、ABC123 各计 1；
    - 其它非空白字符（包括中文、中文标点、英文标点、引号、连字符等）逐个计 1；
    - 空白字符不计。

    返回：(character_count, word_count, total)
    其中 total = character_count + word_count。

    说明：旧版本只把部分 CJK 字符/标点放进 character_count，像“ ”、"-"
    这样的字符会被漏掉，导致与 WPS 的正文“字数”相差。
    """
    character_count = 0
    word_count = 0
    in_word = False

    for line in body_lines:
        s = line.strip()
        if not s:
            continue

        for ch in s:
            if ch.isspace():
                in_word = False
                continue

            # 与原程序保持一致：连续 ASCII 字母/数字作为 1 个单词。
            if ch.isascii() and ch.isalnum():
                if not in_word:
                    word_count += 1
                    in_word = True
            else:
                # WPS 统计中，中文字符、中文标点、英文标点、各种引号、
                # 连字符等非空白字符都计入；不能只限定某几个 Unicode 区段。
                character_count += 1
                in_word = False

    total = character_count + word_count
    return character_count, word_count, total



class ScrollableFrame(ttk.Frame):
    def __init__(self, container, *args, **kwargs):
        super().__init__(container, *args, **kwargs)
        self.canvas = tk.Canvas(self)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.scroll_frame = ttk.Frame(self.canvas)
        self.scroll_frame.grid_columnconfigure(0, weight=1)
        self.scroll_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas_window = self.canvas.create_window((0,0), window=self.scroll_frame, anchor="nw", width=720)
        self.canvas.configure(yscrollcommand=scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        def _on_mousewheel(event):
            self.canvas.yview_scroll(int(-1*(event.delta/120)), "units")
        self.canvas.bind_all("<MouseWheel>", _on_mousewheel)
    def scroll_widget_into_view(self, widget):
        self.canvas.update_idletasks()
        widget_y = widget.winfo_y()
        widget_h = widget.winfo_height()
        view_top = self.canvas.canvasy(0)
        view_bottom = view_top + self.canvas.winfo_height()
        if widget_y < view_top:
            self.canvas.yview_moveto(widget_y / max(self.scroll_frame.winfo_height(),1))
        elif widget_y + widget_h > view_bottom:
            bottom = widget_y + widget_h
            total = max(self.scroll_frame.winfo_height(),1)
            self.canvas.yview_moveto(bottom / total)

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("投稿排版助手 v1.0.14")
        win_w = 780
        screen_h = self.winfo_screenheight()
        win_h = min(620, screen_h - 100)
        self.geometry(f"{win_w}x{win_h}")
        self.minsize(720, 480)
        self.resizable(True, True)
        self.update()
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        x = int((sw - win_w) / 2)
        y = int((sh - win_h) / 2) - 20
        self.geometry(f"{win_w}x{win_h}+{x}+{y}")
        self.cfg = load_config()
        self._grid_rows = []
        self._delete_btns = []
        self._del_to_fields = {}
        self._prev_focus = None
        self.scroll_frame = ScrollableFrame(self)
        self.scroll_frame.pack(fill="both", expand=True, padx=10, pady=8)
        self.main = self.scroll_frame.scroll_frame
        content_wrapper = ttk.Frame(self.main)
        content_wrapper.grid(row=0, column=0, sticky="nsew", padx=20)
        content_wrapper.columnconfigure(1, weight=1)
        row_idx = 0
        ttk.Label(content_wrapper, text="1. 选择文件（支持 .docx、.txt、.md；不支持旧版 .doc，请先另存为 .docx）").grid(row=row_idx, column=0, columnspan=2, sticky="w", pady=(0,4))
        row_idx += 1
        fr1 = ttk.Frame(content_wrapper)
        fr1.grid(row=row_idx, column=0, columnspan=2, sticky="ew", pady=2)
        fr1.columnconfigure(0, weight=1)
        row_idx += 1
        self.file_path_var = tk.StringVar()
        e_file = ttk.Entry(fr1, textvariable=self.file_path_var)
        e_file.pack(side="left", fill="x", expand=True)
        btn_browse1 = ttk.Button(fr1, text="浏览文件 [B]", command=self.select_file, takefocus=True)
        btn_browse1.pack(side="left", padx=5)
        self._grid_rows.append([e_file, btn_browse1])
        ttk.Label(content_wrapper, text="2. 选择输出保存文件夹").grid(row=row_idx, column=0, columnspan=2, sticky="w", pady=(10,4))
        row_idx +=1
        fr2 = ttk.Frame(content_wrapper)
        fr2.grid(row=row_idx, column=0, columnspan=2, sticky="ew", pady=2)
        fr2.columnconfigure(0, weight=1)
        row_idx +=1
        self.save_dir_var = tk.StringVar(value=self.cfg.get("last_save_dir",""))
        e_save = ttk.Entry(fr2, textvariable=self.save_dir_var)
        e_save.pack(side="left", fill="x", expand=True)
        btn_browse2 = ttk.Button(fr2, text="浏览文件夹 [L]", command=self.select_save_folder, takefocus=True)
        btn_browse2.pack(side="left", padx=5)
        self._grid_rows.append([e_save, btn_browse2])
        ttk.Label(content_wrapper, text="3. 文章信息").grid(row=row_idx, column=0, columnspan=2, sticky="w", pady=(10,6))
        row_idx +=1
        fr3 = ttk.Frame(content_wrapper)
        fr3.grid(row=row_idx, column=0, columnspan=2, sticky="ew")
        fr3.columnconfigure(1, weight=1)
        row_idx +=1
        ttk.Label(fr3, text="标题：").grid(row=0,column=0,sticky="w")
        self.title_var = tk.StringVar()
        e_title = ttk.Entry(fr3, textvariable=self.title_var)
        e_title.grid(row=0,column=1,sticky="ew",padx=5)
        ttk.Label(fr3, text="作者：").grid(row=1,column=0,sticky="w",pady=3)
        self.author_var = tk.StringVar()
        e_author = ttk.Entry(fr3, textvariable=self.author_var)
        e_author.grid(row=1,column=1,sticky="ew",padx=5)
        self._grid_rows.append([e_title])
        self._grid_rows.append([e_author])
        ttk.Label(content_wrapper, text="4. 联系信息").grid(row=row_idx, column=0, columnspan=2, sticky="w", pady=(10,6))
        row_idx +=1
        fr4 = ttk.Frame(content_wrapper)
        fr4.grid(row=row_idx, column=0, columnspan=2, sticky="ew")
        fr4.columnconfigure(1, weight=1)
        row_idx +=1
        self.contact_vars = {}
        contact_keys = ["姓名","微信","电话","地址","邮编","身份证号","邮箱"]
        for idx, key in enumerate(contact_keys):
            ttk.Label(fr4, text=f"{key}：").grid(row=idx,column=0,sticky="w")
            var = tk.StringVar(value=self.cfg["contact"].get(key,""))
            entry = ttk.Entry(fr4, textvariable=var)
            entry.grid(row=idx,column=1,sticky="ew",padx=5,pady=1)
            self.contact_vars[key] = (var, entry)
            self._grid_rows.append([entry])
        ttk.Label(content_wrapper, text="5. 银行信息（最多4组；如收稿方未要求填写联行号，可留空）").grid(row=row_idx, column=0, columnspan=2, sticky="w", pady=(10,6))
        row_idx +=1
        self.bank_container = ttk.Frame(content_wrapper)
        self.bank_container.grid(row=row_idx, column=0, columnspan=2, sticky="ew")
        self.bank_container.columnconfigure(0, weight=1)
        row_idx +=1
        self.bank_frames = []
        self.bank_data = []
        self.btn_add_bank = ttk.Button(content_wrapper, text="添加银行 [G]", command=self.add_bank_group, takefocus=True)
        self._grid_rows.append([self.btn_add_bank])
        for b in self.cfg["banks"]:
            self.add_bank_group(init_data=b)
        self.btn_add_bank.grid(row=row_idx, column=0, sticky="w",pady=3)
        row_idx +=1
        self.export_btn = ttk.Button(content_wrapper, text="开始排版并导出 [X]", command=self.do_export, takefocus=True)
        self.export_btn.grid(row=row_idx, column=0, columnspan=2, pady=12)
        self._grid_rows.append([self.export_btn])
        self._bind_keys()

        # Windows 风格面板快捷键：Alt+字母直接执行对应按钮。
        self.bind_all("<Alt-b>", lambda e: (self.select_file(), "break")[1])
        self.bind_all("<Alt-l>", lambda e: (self.select_save_folder(), "break")[1])
        self.bind_all("<Alt-g>", lambda e: (self.add_bank_group(), "break")[1])
        self.bind_all("<Alt-x>", lambda e: (self.export_btn.invoke(), "break")[1])
        self.after(100, lambda: self._focus_widget(btn_browse1))

    def _find_widget_position(self, widget):
        for r, row in enumerate(self._grid_rows):
            for c, w in enumerate(row):
                if w is widget:
                    return r, c
        return None, None

    def _focus_widget(self, widget):
        self._prev_focus = self.focus_get()
        widget.focus_set()
        self.scroll_frame.scroll_widget_into_view(widget)

    def _bind_keys(self):
        def move_up(event=None):
            current = self.focus_get()
            if current in self._delete_btns:
                idx = self._delete_btns.index(current)
                if idx > 0:
                    self._focus_widget(self._delete_btns[idx - 1])
                else:
                    _, email_entry = self.contact_vars["邮箱"]
                    self._focus_widget(email_entry)
                return "break"
            r, c = self._find_widget_position(current)
            if r is None:
                return
            if r > 0:
                target_row = self._grid_rows[r - 1]
                if c < len(target_row):
                    self._focus_widget(target_row[c])
                else:
                    self._focus_widget(target_row[-1])
            else:
                self._focus_widget(self._grid_rows[-1][0])
            return "break"
        def move_down(event=None):
            current = self.focus_get()
            if current in self._delete_btns:
                idx = self._delete_btns.index(current)
                if idx < len(self._delete_btns) - 1:
                    self._focus_widget(self._delete_btns[idx + 1])
                else:
                    self._focus_widget(self.btn_add_bank)
                return "break"
            r, c = self._find_widget_position(current)
            if r is None:
                return
            if r < len(self._grid_rows) - 1:
                target_row = self._grid_rows[r + 1]
                if c < len(target_row):
                    self._focus_widget(target_row[c])
                else:
                    self._focus_widget(target_row[-1])
            else:
                self._focus_widget(self._grid_rows[0][0])
            return "break"
        def move_left(event=None):
            current = self.focus_get()
            if current is self.btn_add_bank:
                if self._delete_btns:
                    self._focus_widget(self._delete_btns[-1])
                return "break"
            if current in self._delete_btns:
                entry_card, entry_lxh = self._del_to_fields[current]
                prev = self._prev_focus
                if prev is entry_card:
                    self._focus_widget(entry_card)
                elif prev is entry_lxh:
                    self._focus_widget(entry_lxh)
                elif prev in self._delete_btns:
                    prev_idx = self._delete_btns.index(prev)
                    cur_idx = self._delete_btns.index(current)
                    if prev_idx < cur_idx:
                        self._focus_widget(entry_card)
                    else:
                        self._focus_widget(entry_lxh)
                else:
                    self._focus_widget(entry_card)
                return "break"
            r, c = self._find_widget_position(current)
            if r is None:
                return
            row = self._grid_rows[r]
            if c > 0:
                self._focus_widget(row[c - 1])
            return "break"
        def move_right(event=None):
            current = self.focus_get()
            r, c = self._find_widget_position(current)
            if r is not None and c is not None:
                row = self._grid_rows[r]
                if len(row) == 2 and c == 1 and current not in self._delete_btns and current is not self.btn_add_bank and current is not self.export_btn:
                    for btn, (ec, el) in self._del_to_fields.items():
                        if el is current:
                            self._focus_widget(btn)
                            return "break"
            if r is None:
                return
            row = self._grid_rows[r]
            if c < len(row) - 1:
                self._focus_widget(row[c + 1])
            return "break"
        # 使用 bind_all，让焦点在 Entry 等子控件时方向键也能交给排版助手的导航逻辑。
        # 仅当当前焦点确实属于主窗口时才处理，避免干扰其它弹窗。
        def in_main_window():
            current = self.focus_get()
            if current is None:
                return False
            try:
                return current.winfo_toplevel() is self
            except tk.TclError:
                return False

        def guarded(handler):
            def _handler(event=None):
                if not in_main_window():
                    return
                return handler(event)
            return _handler

        self.bind_all("<Up>", guarded(move_up), add="+")
        self.bind_all("<Down>", guarded(move_down), add="+")
        self.bind_all("<Left>", guarded(move_left), add="+")
        self.bind_all("<Right>", guarded(move_right), add="+")
        def on_return(event=None):
            widget = self.focus_get()
            if isinstance(widget, ttk.Button):
                widget.invoke()
                return "break"
            return "break"
        self.bind_all("<Return>", guarded(on_return), add="+")
        self.bind_all("<KP_Enter>", guarded(on_return), add="+")

    def select_file(self):
        init_dir = self.cfg.get("last_file_dir","")
        fp = filedialog.askopenfilename(
            parent=self,
            initialdir=init_dir,
            filetypes=[("支持文件","*.docx;*.txt;*.md"),("Word文档","*.docx"),("文本文件","*.txt;*.md")]
        )
        if not fp:
            return
        if fp.lower().endswith(".doc"):
            show_alert(self, "格式不支持", "请用 Word 另存为 docx 或 txt 格式", "确定 [O]")
            return
        self.file_path_var.set(fp)
        self.cfg["last_file_dir"] = os.path.dirname(fp)
        title, author = self.parse_file(fp)
        display_title, _ = normalize_title_for_export(title)
        self.title_var.set(display_title)
        self.author_var.set(author)

    def parse_file(self, filepath):
        ext = filepath.lower().split(".")[-1]
        title, author = "", ""
        try:
            if ext in ["txt","md"]:
                with open(filepath,"r",encoding="utf-8") as f:
                    lines = [line.rstrip("\n") for line in f if line.strip()]
                if len(lines)>=1: title = lines[0].strip()
                if len(lines)>=2: author = lines[1].strip()
            elif ext == "docx":
                from docx import Document
                doc = Document(filepath)
                paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                if len(paras)>=1: title = paras[0]
                if len(paras)>=2: author = paras[1]
        except Exception as e:
            show_alert(self, "读取失败", str(e))
        return title, author

    def select_save_folder(self):
        init_dir = self.cfg.get("last_save_dir","")
        folder = filedialog.askdirectory(parent=self, initialdir=init_dir)
        if folder:
            self.save_dir_var.set(folder)
            self.cfg["last_save_dir"] = folder

    def add_bank_group(self, init_data=None):
        if len(self.bank_frames)>=4:
            show_alert(self, "提示", "最多只能添加4组银行信息")
            return
        data = init_data or {"银行名称":"","账号":"","开户支行":"","联行号":""}
        fr = ttk.Frame(self.bank_container)
        fr.pack(fill="x", pady=4)
        fr.columnconfigure(1, weight=1)
        fr.columnconfigure(3, weight=1)
        vars_dict = {}
        ttk.Label(fr, text="银行名称:").grid(row=0, column=0, sticky="w")
        var_bank = tk.StringVar(value=data["银行名称"])
        entry_bank = ttk.Entry(fr, textvariable=var_bank)
        entry_bank.grid(row=0, column=1, padx=3, sticky="ew")
        ttk.Label(fr, text="账号:").grid(row=0, column=2, sticky="w", padx=(8,0))
        var_card = tk.StringVar(value=data["账号"])
        entry_card = ttk.Entry(fr, textvariable=var_card)
        entry_card.grid(row=0, column=3, padx=3, sticky="ew")
        ttk.Label(fr, text="开户支行:").grid(row=1, column=0, sticky="w")
        var_branch = tk.StringVar(value=data["开户支行"])
        entry_branch = ttk.Entry(fr, textvariable=var_branch)
        entry_branch.grid(row=1, column=1, padx=3, sticky="ew")
        ttk.Label(fr, text="联行号:").grid(row=1, column=2, sticky="w", padx=(8,0))
        var_lxh = tk.StringVar(value=data["联行号"])
        entry_lxh = ttk.Entry(fr, textvariable=var_lxh)
        entry_lxh.grid(row=1, column=3, padx=3, sticky="ew")
        vars_dict["银行名称"] = (var_bank, entry_bank)
        vars_dict["账号"] = (var_card, entry_card)
        vars_dict["开户支行"] = (var_branch, entry_branch)
        vars_dict["联行号"] = (var_lxh, entry_lxh)
        btn_del = ttk.Button(fr, text="删除本组 [D]", takefocus=True)
        btn_del.grid(row=0, column=4, rowspan=2, padx=4)

        # Alt+D 属于“这一组银行信息”的快捷键，而不是“删除按钮”的快捷键。
        # 因此焦点停在本组任意输入框时，都可以直接删除整组。
        def delete_group_shortcut(event=None):
            btn_del.invoke()
            return "break"

        for entry_widget in (entry_bank, entry_card, entry_branch, entry_lxh, btn_del):
            entry_widget.bind("<Alt-KeyPress-d>", delete_group_shortcut)
        try:
            insert_idx = self._grid_rows.index([self.btn_add_bank])
        except (ValueError, AttributeError):
            insert_idx = len(self._grid_rows)
        row_a = [entry_bank, entry_card, btn_del]
        row_b = [entry_branch, entry_lxh]
        self._grid_rows.insert(insert_idx, row_a)
        insert_idx += 1
        self._grid_rows.insert(insert_idx, row_b)
        new_rows = [row_a, row_b]
        self._delete_btns.append(btn_del)
        self._del_to_fields[btn_del] = (entry_card, entry_lxh)
        def del_this():
            del_row_idx = None
            for r, row in enumerate(self._grid_rows):
                if row is row_a:
                    del_row_idx = r
                    break
            for row in new_rows:
                try:
                    self._grid_rows.remove(row)
                except ValueError:
                    pass
            try:
                self._delete_btns.remove(btn_del)
            except ValueError:
                pass
            if btn_del in self._del_to_fields:
                del self._del_to_fields[btn_del]
            idx = self.bank_frames.index(fr)
            self.bank_frames.pop(idx)
            self.bank_data.pop(idx)
            fr.destroy()
            if del_row_idx is not None and del_row_idx < len(self._grid_rows):
                self._focus_widget(self._grid_rows[del_row_idx][0])
            elif len(self._grid_rows) > 0:
                self._focus_widget(self._grid_rows[-1][0])
        btn_del.config(command=del_this)
        self.bank_frames.append(fr)
        self.bank_data.append(vars_dict)
        if init_data is None:
            self._focus_widget(entry_bank)

    def get_bank_list(self):
        res = []
        for b in self.bank_data:
            d = {}
            for k in ["银行名称","账号","开户支行","联行号"]:
                var, entry = b[k]
                d[k] = var.get().strip()
            res.append(d)
        return res

    def open_folder(self, folder_path):
        path = str(folder_path)
        if platform.system() == "Windows":
            os.startfile(path)
        elif platform.system() == "Darwin":
            os.system(f"open '{path}'")
        else:
            os.system(f"xdg-open '{path}'")

    def do_export(self):
        for key, (var,entry) in self.contact_vars.items():
            entry.config(background="white")
        for bank in self.bank_data:
            for k in ["银行名称","账号","开户支行","联行号"]:
                var,entry = bank[k]
                entry.config(background="white")
        title = self.title_var.get().strip()
        author = self.author_var.get().strip()

        display_title, file_title = normalize_title_for_export(title)
        save_root = self.save_dir_var.get().strip()
        if not all([title, author, save_root]):
            show_alert(self, "缺失", "标题、作者、保存文件夹不能为空")
            return
        error_text = ""
        contact = {}
        for key, (var,entry) in self.contact_vars.items():
            val = var.get().strip()
            contact[key] = val
            if not val:
                entry.config(background="#ffcccc")
                error_text += f"联系信息：{key} 为空\n"
        bank_list = self.get_bank_list()
        for bank_idx, bank in enumerate(bank_list):
            for ck in ["银行名称","账号","开户支行"]:
                val = bank[ck]
                var,entry = self.bank_data[bank_idx][ck]
                if not val:
                    entry.config(background="#ffcccc")
                    error_text += f"第{bank_idx+1}组银行信息：{ck} 为空\n"
        if error_text:
            show_alert(self, "必填项缺失", error_text, width=620, height=300)
            return

        sub_folder_name = f"《{file_title}》投稿"
        out_folder = Path(save_root) / sub_folder_name
        fn_base = f"《{file_title}》{author}"

        if out_folder.exists():
            choice = ask_export_conflict(self, out_folder)
            if choice == "cancel":
                return
            elif choice == "saveas":
                # 可以反复输入：名称为空或目标已存在时，提示后重新打开输入框。
                while True:
                    new_folder_name = ask_saveas_folder_name(self, sub_folder_name)
                    if new_folder_name is None:
                        return

                    new_folder_name = new_folder_name.strip()
                    if not new_folder_name:
                        show_alert(
                            self,
                            "提示",
                            "文件夹名称不能为空，请重新输入。"
                        )
                        continue

                    out_folder = Path(save_root) / new_folder_name

                    if out_folder.exists():
                        show_alert(
                            self,
                            "文件夹已存在",
                            f"文件夹：\n{out_folder}\n已经存在，请换一个名称。"
                        )
                        continue

                    break

            elif choice == "overwrite":
                docx_path = out_folder / f"{fn_base}.docx"
                txt_path = out_folder / f"{fn_base}.txt"

                del_ok = False
                while not del_ok:
                    try:
                        for output_path in (docx_path, txt_path):
                            if output_path.exists():
                                output_path.unlink()
                        del_ok = True
                    except OSError:
                        ans = ask_retry_cancel(
                            self,
                            "文件占用",
                            "要覆盖的同名文件正在被占用，请关闭相关文件后重试"
                        )
                        if not ans:
                            return

        if not out_folder.exists():
            try:
                out_folder.mkdir(parents=True, exist_ok=False)
            except OSError as e:
                show_alert(
                    self,
                    "无法创建文件夹",
                    f"文件夹无法创建：\n\n{out_folder.name}\n\n"
                    f"错误信息：{e}",
                    width=620,
                    height=280
                )
                return

        docx_path = out_folder / f"{fn_base}.docx"
        txt_path = out_folder / f"{fn_base}.txt"
        fp = self.file_path_var.get()
        content_lines = self.read_full_text(fp)
        raw_body_lines = content_lines[2:] if len(content_lines)>=3 else []

        # 仅正文做半角标点转全角
        processed_body_lines = []
        for ln in raw_body_lines:
            new_ln = convert_halfwidth_punct(ln)
            processed_body_lines.append(new_ln)

        cn_cnt, word_cnt, total_cnt = count_cn_and_word(processed_body_lines)

        doc = Document()
        p = doc.add_paragraph()
        _set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(display_title)
        _set_run_font(run, 14)
        run.bold = False
        p = doc.add_paragraph()
        _set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(author)
        _set_run_font(run, 10.5)

        for line in processed_body_lines:
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph()
            _set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=Pt(21))
            run = p.add_run(line)
            _set_run_font(run, 10.5)

        p = doc.add_paragraph()
        _set_para(p, align=WD_ALIGN_PARAGRAPH.RIGHT)
        run = p.add_run(f"（正文：{total_cnt} 字）")
        _set_run_font(run, 10.5)
        doc.add_paragraph()
        p = doc.add_paragraph()
        _set_para(p)
        run = p.add_run("联系信息：")
        _set_run_font(run, 10.5)
        for key in ["姓名","微信","电话","地址","邮编","身份证号","邮箱"]:
            val = contact.get(key,"")
            p = doc.add_paragraph()
            _set_para(p)
            run = p.add_run(f"{key}：{val}")
            _set_run_font(run, 10.5)
        doc.add_paragraph()
        for idx,bank in enumerate(bank_list):
            p = doc.add_paragraph()
            _set_para(p)
            run = p.add_run(bank["银行名称"])
            _set_run_font(run, 10.5)
            for k in ["账号","开户支行"]:
                val = bank[k]
                p = doc.add_paragraph()
                _set_para(p)
                run = p.add_run(f"{k}：{val}")
                _set_run_font(run, 10.5)
            lxh_val = bank["联行号"].strip()
            if lxh_val:
                p = doc.add_paragraph()
                _set_para(p)
                run = p.add_run(f"联行号：{lxh_val}")
                _set_run_font(run, 10.5)
            if idx != len(bank_list)-1:
                doc.add_paragraph()
        doc.add_paragraph()

        save_ok = False
        while not save_ok:
            try:
                doc.save(docx_path)
                save_ok = True
            except PermissionError:
                res = ask_retry_cancel(
                    self,
                    "文件占用",
                    "目标文件或文件夹正在被占用，请关闭相关文件/文件夹后重试"
                )
                if not res:
                    return

        txt_content = []
        txt_content.append(display_title)
        txt_content.append(author)
        for bl in processed_body_lines:
            bl = bl.strip()
            if bl:
                txt_content.append("　　"+bl)
        txt_content.append(f"（正文：{total_cnt} 字）")
        txt_content.append("")
        txt_content.append("联系信息：")
        for k in ["姓名","微信","电话","地址","邮编","身份证号","邮箱"]:
            txt_content.append(f"{k}：{contact.get(k,'')}")
        txt_content.append("")
        for idx,bank in enumerate(bank_list):
            txt_content.append(bank["银行名称"])
            txt_content.append(f"账号：{bank['账号']}")
            txt_content.append(f"开户支行：{bank['开户支行']}")
            lxh_val = bank["联行号"].strip()
            if lxh_val:
                txt_content.append(f"联行号：{lxh_val}")
            if idx != len(bank_list)-1:
                txt_content.append("")
        txt_content.append("")
        with open(txt_path,"w",encoding="utf-8") as f:
            f.write("\n".join(txt_content))
        self.cfg["contact"] = {k:v.get().strip() for k,(v,e) in self.contact_vars.items()}
        self.cfg["banks"] = self.get_bank_list()
        save_config(self.cfg)
        win = tk.Toplevel(self)
        win.title("导出成功")
        win.transient(self)
        win.resizable(False, False)
        win.grab_set()
        win.geometry("320x120")
        win.update()
        sw2 = win.winfo_screenwidth()
        sh2 = win.winfo_screenheight()
        win.geometry(f"320x120+{int((sw2-320)/2)}+{int((sh2-120)/2)}")
        ttk.Label(win, text=f"文件已导出到：\n{out_folder}", wraplength=300).pack(pady=15)
        def close_all(event=None):
            try:
                win.grab_release()
            except tk.TclError:
                pass
            if win.winfo_exists():
                win.destroy()
            self.open_folder(out_folder)
            self.quit()
            return "break"
        close_btn = ttk.Button(win, text="关闭 [C]", command=close_all)
        close_btn.pack()
        win.bind("<Alt-KeyPress-c>", close_all)
        win.bind("<Return>", close_all)
        win.bind("<Escape>", close_all)
        win.protocol("WM_DELETE_WINDOW", close_all)
        win.after_idle(lambda: (win.focus_force(), close_btn.focus_set()))
        win.after(3000, close_all)

    def read_full_text(self,filepath):
        ext = filepath.lower().split(".")[-1]
        lines = []
        if ext in ["txt","md"]:
            with open(filepath,"r",encoding="utf-8") as f:
                lines = [line.rstrip("\n") for line in f]
        elif ext == "docx":
            from docx import Document
            doc = Document(filepath)
            lines = [p.text for p in doc.paragraphs]
        return lines

if __name__ == "__main__":
    app = App()
    app.mainloop()
